# backend/engine.py
# Generic calculation engine — works with any Excel + config.json

import uuid
import shutil
import subprocess
import tempfile
from pathlib import Path

import openpyxl

from config import LIBREOFFICE_BIN


def _get_session_dirs():
    base = Path(tempfile.gettempdir()) / "rater_sessions"
    input_dir = base / "input"
    output_dir = base / "output"
    input_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)
    return input_dir, output_dir


def _coerce_by_type(value, value_type):
    if value is None:
        return None
    if value_type != "number":
        return value

    if isinstance(value, (int, float)):
        return value

    text = str(value).strip()
    if text == "":
        return None
    try:
        return float(text) if "." in text else int(text)
    except ValueError:
        return value


def _write_schedule_inputs(ws, config, input_data):
    """
    Write dynamic schedule rows when config contains schedule metadata.

    Returns a set of cell references that were controlled by schedule logic,
    so flat writers can skip those cells and avoid conflicting writes.
    """
    schedule_defs = config.get("schedules") or []
    if not schedule_defs:
        return set()

    schedule_payload = input_data.get("schedules") or input_data.get("_schedules")
    if not isinstance(schedule_payload, dict):
        return set()

    write_rules = config.get("writeRules") or {}
    clear_unused = bool(write_rules.get("clearUnusedRows", True))

    controlled_cells = set()

    for sched in schedule_defs:
        key = sched.get("key")
        if not key:
            continue

        row_start = sched.get("rowStart")
        row_end = sched.get("rowEnd")
        columns = sched.get("columns") or []
        if not isinstance(row_start, int) or not isinstance(row_end, int):
            continue
        if row_end < row_start:
            continue

        rows_data = schedule_payload.get(key) or []
        if not isinstance(rows_data, list):
            rows_data = []

        for idx, row_num in enumerate(range(row_start, row_end + 1)):
            row_data = rows_data[idx] if idx < len(rows_data) and isinstance(rows_data[idx], dict) else {}
            row_active = any(v not in (None, "") for v in row_data.values())

            for col_def in columns:
                col = col_def.get("column")
                field = col_def.get("field")
                value_type = col_def.get("type")
                if not col or not field:
                    continue

                cell_ref = f"{col}{row_num}"
                controlled_cells.add(cell_ref)

                if row_active and field in row_data:
                    ws[cell_ref] = _coerce_by_type(row_data.get(field), value_type)
                elif clear_unused:
                    ws[cell_ref] = None

    return controlled_cells


def _write_cell(wb, default_sheet_name, cell_ref, value):
    if "!" in cell_ref:
        sheet_part, coord = cell_ref.split("!", 1)
        sheet_part = sheet_part.strip("'\"")
        wb[sheet_part][coord] = value
    else:
        wb[default_sheet_name][cell_ref] = value

def _read_cell(wb, default_sheet_name, cell_ref):
    if "!" in cell_ref:
        sheet_part, coord = cell_ref.split("!", 1)
        sheet_part = sheet_part.strip("'\"")
        return wb[sheet_part][coord].value
    return wb[default_sheet_name][cell_ref].value

def calculate(template_path: Path, config: dict, input_data: dict, keep_file: bool = False) -> dict:
    """
    Generic calculation cycle.

    Args:
        template_path: path to the master .xlsx template (never modified)
        config: dict with "sheet", "inputs", "outputs" keys
        input_data: dict of { field_name: value } from the user form
        keep_file: if True, keep the output .xlsx and return its path

    Returns:
        dict of output values keyed by field name
    """
    session_id = str(uuid.uuid4())
    input_dir, output_dir = _get_session_dirs()
    session_file = input_dir / f"{session_id}.xlsx"
    output_file = output_dir / f"{session_id}.xlsx"

    # Build lookup: field_name -> cell reference
    input_map = {inp["field"]: inp["cell"] for inp in config["inputs"]}
    input_type_map = {inp["field"]: inp.get("type", "text") for inp in config["inputs"]}
    output_map = {out["field"]: out["cell"] for out in config["outputs"]}
    sheet_name = config["sheet"]

    try:
        # Step 1: Copy template
        shutil.copy(template_path, session_file)

        # Step 2: Write inputs
        wb = openpyxl.load_workbook(session_file)
        ws = wb[sheet_name]

        # Optional schedule-mode write path (backward-compatible).
        # If no schedule payload/definitions exist, this is a no-op.
        schedule_cells = _write_schedule_inputs(ws, config, input_data)

        for field, value in input_data.items():
            if field in {"schedules", "_schedules"}:
                continue
            cell_ref = input_map.get(field)
            if cell_ref and cell_ref not in schedule_cells:
                val_type = input_type_map.get(field, "text")
                _write_cell(wb, sheet_name, cell_ref, _coerce_by_type(value, val_type))

        wb.save(session_file)
        wb.close()

        # Step 3: LibreOffice recalculation
        result = subprocess.run(
            [
                LIBREOFFICE_BIN,
                "--headless",
                "--calc",
                "--convert-to", "xlsx",
                "--outdir", str(output_dir),
                str(session_file),
            ],
            capture_output=True,
            text=True,
            timeout=60,
        )

        if result.returncode != 0:
            raise RuntimeError(
                f"LibreOffice failed (exit={result.returncode}): "
                f"stderr={result.stderr.strip()}"
            )

        if not output_file.exists():
            raise RuntimeError(
                f"LibreOffice did not produce output file. "
                f"stdout={result.stdout.strip()}"
            )

        # Step 4: Read outputs
        wb_out = openpyxl.load_workbook(output_file, data_only=True)

        outputs = {}
        for field, cell_ref in output_map.items():
            val = _read_cell(wb_out, sheet_name, cell_ref)
            if isinstance(val, float):
                val = round(val, 4)
            outputs[field] = val

        wb_out.close()

        if keep_file:
            outputs["_output_file"] = str(output_file)

        return outputs

    finally:
        # Step 5: Cleanup
        if session_file.exists():
            try:
                session_file.unlink()
            except OSError:
                pass
        if not keep_file and output_file.exists():
            try:
                output_file.unlink()
            except OSError:
                pass
